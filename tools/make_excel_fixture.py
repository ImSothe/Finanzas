# -*- coding: utf-8 -*-
"""Genera tests/fixtures/excel_fixture.json a partir de la plantilla REAL de Excel (con datos inventados).

Rellena la pestaña Piloto con una configuración sintética, deja que `formulas` calcule la herencia mes a mes y exporta la
cuadrícula de valores. Así las pruebas de Code.gs leen exactamente lo que leería Google Sheets.

Uso: python tools/make_excel_fixture.py  ruta/a/Finanzas_Pro_AntiPereza.xlsx
"""
import json
import os
import sys
import warnings
import openpyxl
import formulas

warnings.filterwarnings('ignore')
SRC = sys.argv[1]
TMP = 'seed/_fixture_tmp.xlsx'
S_P, S_G, S_X = '⚙️ Piloto Automático', '📝 Gastos', '🚀 Extras'

wb = openpyxl.load_workbook(SRC)
p = wb[S_P]
# --- configuración sintética (cifras inventadas)
tot = next(r for r in range(1, 40) if str(p.cell(r, 2).value).startswith('TOTAL GASTOS FIJOS'))
year0 = int(p.cell(next(r for r in range(40, 90) if str(p.cell(r, 2).value).startswith('Primer año')), 3).value)
p['D9'] = 2000                                  # sueldo desde enero del primer año
p.cell(9, 4 + 12).value = 2200                  # sube en enero del año siguiente
p['D10'] = 600                                  # vivienda fija
hdr_sub = next(r for r in range(tot, 90) if str(p.cell(r, 2).value) == 'Nombre')
p.cell(hdr_sub + 1, 2).value = 'Streaming'; p.cell(hdr_sub + 1, 3).value = 10
p.cell(hdr_sub + 2, 2).value = 'Otra'; p.cell(hdr_sub + 2, 3).value = 5; p.cell(hdr_sub + 2, 5).value = 'No'
wb.save(TMP)

sol = formulas.ExcelModel().loads(TMP).finish().calculate()
base = "'[%s]%s'!" % (os.path.basename(TMP), S_P.upper())
grid = []
for r in range(1, 101):
    row = []
    for c in range(1, 100):
        raw = p.cell(r, c).value
        ref = openpyxl.utils.get_column_letter(c) + str(r)
        if isinstance(raw, str) and raw.startswith('='):
            v = sol.get(base + ref)
            val = v.value[0, 0] if v is not None else ''
            val = '' if str(val) in ('', 'None', '#N/A') else (float(val) if isinstance(val, (int, float)) or hasattr(val, 'item') else val)
            if isinstance(val, str) and val.startswith('#'):
                val = ''
        else:
            val = raw if raw is not None else ''
        if hasattr(val, 'isoformat'):
            val = val.isoformat()
        row.append(val)
    grid.append(row)


def cabeceras(name, hasta=9, cols=14):
    ws = wb[name]
    return [[(ws.cell(r, c).value if ws.cell(r, c).value is not None else '') for c in range(1, cols + 1)] for r in range(1, hasta + 1)]


out = {'year0': year0, 'piloto': grid, 'gastosCab': cabeceras(S_G), 'extrasCab': cabeceras(S_X, 9, 11)}
os.makedirs('tests/fixtures', exist_ok=True)
json.dump(out, open('tests/fixtures/excel_fixture.json', 'w', encoding='utf-8'), ensure_ascii=False)
os.remove(TMP)
print('fixture escrito:', len(grid), 'filas x', len(grid[0]), 'columnas; year0 =', year0,
      '| sueldo ene/primer año =', grid[8][3], '| ene/siguiente =', grid[8][15], '| vivienda =', grid[9][3], grid[9][50])
